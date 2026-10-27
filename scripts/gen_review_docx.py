#!/usr/bin/env python3
"""Generate the Farmsky cross-app architecture & security review as a Word doc."""
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

BRAND = RGBColor(0x0F, 0x76, 0x6E)      # teal
RED = RGBColor(0xC0, 0x1A, 0x1A)
ORANGE = RGBColor(0xB4, 0x53, 0x09)
YELLOW = RGBColor(0x8A, 0x6D, 0x00)
GREY = RGBColor(0x55, 0x55, 0x55)

doc = Document()

# Base styles
style = doc.styles['Normal']
style.font.name = 'Calibri'
style.font.size = Pt(10.5)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:val'), 'clear'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def sev_run(p, sev):
    colors = {'High': RED, 'Medium': ORANGE, 'Low': YELLOW}
    r = p.add_run(f'[{sev}] ')
    r.bold = True; r.font.color.rgb = colors.get(sev, GREY)

def h(text, level=1, color=BRAND):
    p = doc.add_heading(text, level=level)
    for r in p.runs:
        r.font.color.rgb = color
    return p

def para(text, bold=False, italic=False, size=10.5, color=None, space_after=6):
    p = doc.add_paragraph()
    r = p.add_run(text); r.bold = bold; r.italic = italic; r.font.size = Pt(size)
    if color: r.font.color.rgb = color
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text, sev=None):
    p = doc.add_paragraph(style='List Bullet')
    if sev: sev_run(p, sev)
    p.add_run(text)
    return p

def make_table(headers, rows, widths=None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Light Grid Accent 1'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr = t.rows[0].cells
    for i, hh in enumerate(headers):
        hdr[i].text = ''
        run = hdr[i].paragraphs[0].add_run(hh); run.bold = True; run.font.size = Pt(9.5)
        run.font.color.rgb = RGBColor(0xFF,0xFF,0xFF)
        shade(hdr[i], '0F766E')
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].text = ''
            run = cells[i].paragraphs[0].add_run(str(val)); run.font.size = Pt(9)
    if widths:
        for row in t.rows:
            for i, w in enumerate(widths):
                row.cells[i].width = Inches(w)
    return t

# ---------------- COVER ----------------
title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = title.add_run('Farmsky Multi-App System\nArchitecture & Security Review')
r.bold = True; r.font.size = Pt(24); r.font.color.rgb = BRAND
sub = doc.add_paragraph(); sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
rs = sub.add_run('Duplicate / Conflicting Functionality Audit + Security Gap Assessment with Recommendations')
rs.italic = True; rs.font.size = Pt(12); rs.font.color.rgb = GREY
meta = doc.add_paragraph(); meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
meta.add_run('Scope: Equipment (farm-sky-webapp) · Feed (feed) · Nia (input)\n').font.size = Pt(10)
meta.add_run('Prepared for: Farmsky Ventures Ops  ·  Document type: Findings & Remediation Backlog').font.size = Pt(10)
doc.add_paragraph()

leg = doc.add_paragraph(); leg.alignment = WD_ALIGN_PARAGRAPH.CENTER
leg.add_run('Severity legend:  ').bold = True
for lbl, col in [('High', RED), ('Medium', ORANGE), ('Low', YELLOW)]:
    rr = leg.add_run(f'{lbl}   '); rr.bold = True; rr.font.color.rgb = col

doc.add_page_break()

# ---------------- 0. EXEC SUMMARY ----------------
h('0. Executive Summary', 1)
para('The three apps began as a single codebase and were forked into a central host '
     '(Equipment) plus two satellites (Feed, Nia). Over successive feature rounds the forks '
     'drifted: the same capability is now implemented under different file names, with different '
     'internal sizes, and occasionally with subtly different behaviour. This review flags:')
bullet('9 duplicate / parallel / redundant mechanisms (architectural redundancy, state-desync risk, maintenance cost).')
bullet('4 conflicting / confusing / opposite-function patterns (ambiguous source of truth, dual code paths).')
bullet('8 security gaps / risks / exposures (High to Low severity).')
para('Each finding carries a concrete recommendation. Section 5 proposes a phased consolidation '
     'plan that reduces complexity while preserving zero-regression guarantees.', italic=True)

# ---------------- 1. TOPOLOGY ----------------
h('1. System Topology (as-built)', 1)
make_table(
    ['App', 'Repo', 'Role', 'Payment posture', 'Lines (index / app.js)'],
    [
        ['Equipment', 'farm-sky-webapp', 'Central HOST + own marketplace', 'Hosts gateway; also a tenant', '8356 / 8091'],
        ['Feed', 'feed', 'Livestock satellite (tenant)', 'Routes to Equipment gateway', '8297 / 8083'],
        ['Nia', 'input', 'Crop/horticulture satellite', 'Routes to Equipment gateway', '6440 / 5712'],
    ],
    widths=[0.9, 1.5, 1.8, 1.7, 1.3])
para('')
para('All three share: Hono + TypeScript backend, dual build (dist-node/server.js authoritative '
     '+ dist/_worker.js for Cloudflare), PostgreSQL with FORCE RLS, feature_flags lifecycle '
     '(sandbox -> testing -> live), and a vanilla-JS SPA (frontend/static/app.js).')

# ---------------- 2. DUPLICATES ----------------
h('2. Duplicate / Parallel / Redundant Functionalities', 1)

h('2.1 Payment-gateway module: three divergent copies of the same host', 2, RED)
para('The central gateway logic exists in differently-named, differently-sized files in each repo '
     'despite implementing the same functions (candidateSecrets, genRef, loadClient, loadMarketplace, '
     'setTenantScope, verifyIpn, settleCallback, notifyOriginApp, ...):')
make_table(['App', 'Host-gateway file', 'Size', 'Shared helper', 'Client file'],
    [['Equipment', 'payment-gateway.ts', '72 KB', 'payments-shared.ts', '(none)'],
     ['Feed', 'payment-gateway-host.ts', '68 KB', 'payment-gateway-shared.ts', 'payment-gateway-client.ts'],
     ['Nia', 'payment-gateway.ts', '47 KB', 'payments-shared.ts', 'payments-client.ts']],
    widths=[1.0, 1.8, 0.7, 1.7, 1.7])
para('')
p = doc.add_paragraph(); sev_run(p, 'High')
p.add_run('Risk: ').bold = True
p.add_run('High architectural redundancy. A fix to IPN verification or settlement must be hand-ported '
          'to 3 files with 3 different names and 3 different internal states -> high risk of STATE '
          'DESYNCHRONIZATION. Satellites arguably should not carry a 47-68 KB host gateway at all.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Extract ONE canonical @farmsky/payment-gateway (host) + ONE @farmsky/payment-client (tenant) '
          'package. Satellites import the client only. If a monorepo is not viable short-term, adopt '
          'identical file names enforced by a CI diff gate.')

h('2.2 wallet-gateway.ts re-implements gateway primitives', 2, RED)
para('wallet-gateway.ts (Equipment & Feed, ~24 KB; ABSENT in Nia) redefines genRef, loadClient, '
     'candidateSecrets, auditSecurity and authenticate — functions that already exist in '
     'payment-gateway.ts. Two independent copies of the HMAC auth path now exist in the same repo.')
p = doc.add_paragraph(); sev_run(p, 'High'); p.add_run('Risk: ').bold = True
p.add_run('An auth/signing change applied to the payment gateway but not the wallet gateway creates an '
          'auth-logic split-brain. Nia lacking the file means wallet features silently differ by app.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Factor shared auth/signing primitives into one internal module imported by both gateways; '
          'make Nia wallet posture explicit (add module, or document wallets unsupported on Nia).')

h('2.3 Dual checkout endpoint aliases (/api/murabaha/* vs /api/checkout/*)', 2, ORANGE)
para('Both legacy /api/murabaha/{quote,apply,apply-bundle} and newer /api/checkout/{quote,apply,'
     'quote-bundle,kyc-check} are registered to the same handlers — but the aliasing is inconsistent: '
     'Equipment & Feed register BOTH /api/checkout/apply and /api/murabaha/apply; Nia registers ONLY '
     '/api/murabaha/apply (no /api/checkout/apply alias).')
p = doc.add_paragraph(); sev_run(p, 'Medium'); p.add_run('Risk: ').bold = True
p.add_run('A client coded against /api/checkout/apply works on Equipment/Feed but 404s on Nia. Two '
          'public names for one behaviour doubles documentation & test surface.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Pick ONE canonical namespace (/api/checkout/*), make legacy names thin redirects (or register '
          'the identical alias set in all three), schedule legacy for deprecation.')

h('2.4 financingQuote() has drifted across apps', 2, ORANGE)
para('The single most important money function differs by app: Equipment 146 lines (tenure + interest + '
     'PAYGO), Feed 106 lines (Murabaha, charge-free), Nia 85 lines (interest, no tenure). Each was edited '
     'independently during the cart/checkout overhaul.')
p = doc.add_paragraph(); sev_run(p, 'Medium'); p.add_run('Risk: ').bold = True
p.add_run('The "one source of truth for price" is actually THREE sources. A rounding/deposit rule corrected '
          'in one app can silently differ in another, re-introducing the cart-vs-contract bug class already '
          'fixed once.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Promote financingQuote + computeProcessingFee + requiredNum into a shared pricing module with a '
          'shared unit-test vector set run in all three CIs. Per-app policy becomes a config flag, not a '
          'forked function body.')

h('2.5 buni.ts (KCB Bank gateway) drift', 2, ORANGE)
para('buni.ts is 367 lines in Equipment & Feed but only 91 lines in Nia — Nia carries a stripped/older KCB '
     'integration.')
p = doc.add_paragraph(); sev_run(p, 'Medium'); p.add_run('Risk: ').bold = True
p.add_run('KCB payments behave/fail differently on Nia vs others; a KCB bugfix lands in 2 of 3 apps.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Converge to the 367-line canonical buni.ts (or the shared package from 2.1); delete the divergent copy.')

h('2.6 Parallel configuration mechanisms (env-var vs DB app_settings)', 2, ORANGE)
para('Configuration is read from TWO parallel stores: 547 c.env.* reads (deployment env vars) AND a DB '
     'app_settings table via getSetting/setSetting. Several values (processing fee, feature availability, '
     'payment creds) can be influenced by BOTH, with implicit precedence.')
p = doc.add_paragraph(); sev_run(p, 'Medium'); p.add_run('Risk: ').bold = True
p.add_run('Exactly the "parallel configuration mechanism -> high redundancy + state desync" pattern. An '
          'operator changing a DB setting may be silently overridden by an env var (or vice-versa) with no '
          'visible precedence.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Define a single config-resolution contract (e.g. env var is the hard floor; DB setting may narrow '
          'but never widen; precedence logged). Centralize reads behind one resolveConfig(key) helper and '
          'document the precedence table.')

h('2.7 Processing-fee computed/normalized in multiple spots', 2, YELLOW)
para('DEFAULT_PROCESSING_FEE, normalizeProcessingFee, computeProcessingFee plus inline '
     'getSetting(c,"processing_fee",...) calls appear at several call sites per app (and in all 3). '
     'Recommendation: single getProcessingFeeConfig(c) accessor (subsumed by 2.4\'s shared module).')

h('2.8 Inconsistent backend/middleware/ presence', 2, YELLOW)
para('Equipment and Nia have backend/middleware/auth.ts; Feed has no middleware directory (auth is inline). '
     'Recommendation: standardize the middleware location across all three.')

h('2.9 RLS-context helpers: withAdminContext (pooled) vs runAsAdmin (pinned)', 2, YELLOW)
para('Two mechanisms set the admin DB context: the older pooled withAdminContext (sets a GUC then issues a '
     'separate pooled query — racy under load) and the newer runAsAdmin (pins one connection — correct). Both '
     'remain in use (~107 context-related lines in Equipment alone).')
p = doc.add_paragraph(); sev_run(p, 'Low'); p.add_run('Risk: ').bold = True
p.add_run('The pooled variant can route the follow-up query to a connection that never saw the GUC -> '
          'intermittent RLS denials. Two patterns invite using the wrong one.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Migrate all privileged writes to runAsAdmin; keep withAdminContext only for benign read paths or '
          'replace entirely. Add a lint rule forbidding pool.query inside a withAdminContext block.')

# ---------------- 3. CONFLICTS ----------------
h('3. Conflicting / Confusing / Opposite-Function Patterns', 1)

h('3.1 Cross-marketplace checkout is "always CASH" while in-app financing exists', 2, ORANGE)
para('Feed\'s cross-marketplace path (ref XMP) forces cash only ("financing is scoped to the owning '
     'platform"), yet the same product in its home app may be financable. The SAME product can be cash-only '
     'or financable depending on which app you check out from — an opposite/confusing behaviour.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Make the rule explicit in the UI ("cross-marketplace items are cash-only"), OR route '
          'cross-marketplace financing to the owning app\'s quote engine so terms stay consistent (preferred; '
          'aligns with the cross-marketplace pricing-origin rule already adopted in checkout).')

h('3.2 Two KYC gates with different trigger logic', 2, ORANGE)
para('KYC is enforced in checkoutApplyHandler, again in apply-bundle, and surfaced via '
     '/api/checkout/kyc-check. The "financing requires verified KYC; cash bypasses KYC" rule is '
     're-implemented at each site (>=3 places per app x 3 apps = 9 edit points).')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('One assertCheckoutAllowed(customer, items) guard used by all checkout entry points.')

h('3.3 Four overlapping app-identity flags', 2, YELLOW)
para('APP_TYPE vs APP_MARKETPLACE vs product app_scope vs product source_platform vs product marketplace — '
     'overlapping notions of "which app / which marketplace" with implicit mapping. Easy to filter on the '
     'wrong dimension (a data-isolation bug waiting to happen). Recommendation: collapse to the minimum set '
     'and add one appIdentity(c) helper. See the reference table in Section 6.')

h('3.4 demo_otp returned in API responses', 2, YELLOW)
para('Several endpoints return demo_otp in JSON when SMS is simulated — confusing in production and a genuine '
     'exposure if simulation mode is ever on in prod. Recommendation: gate demo_otp behind an explicit '
     'DEMO_MODE flag that is hard-off in production builds.')

# ---------------- 4. SECURITY ----------------
h('4. Security Gaps, Risks & Exposures', 1)

h('4.1 CORS reflects any origin WITH credentials', 2, RED)
para('app.use("/api/*", cors({ origin: (origin) => origin || "*", credentials: true, ... })) — reflecting '
     'the caller\'s Origin AND allowing credentials lets any website make authenticated cross-site requests '
     'with the victim\'s cookie session. Present in all three apps.')
p = doc.add_paragraph(); sev_run(p, 'High'); p.add_run('Risk: ').bold = True
p.add_run('CSRF / cross-site credential theft against the cookie-based session.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Replace with an explicit allow-list of known app origins (equipment/feed/nia .farmsky.africa + '
          'sandbox hosts). Never combine wildcard/reflected origin with credentials: true.')

h('4.2 OTP exposure in responses + plaintext OTP storage', 2, ORANGE)
para('demo_otp is returned in simulation mode, and OTP codes are stored in plaintext in otp_codes.code. '
     'Mitigations exist (5-attempt cap, 5-min expiry, single-use) but a DB read exposes live codes.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Hash OTPs at rest (reuse the PBKDF2 util or an HMAC with a server pepper); strip demo_otp unless '
          'DEMO_MODE is explicitly enabled.')

h('4.3 HMAC "fallback secret" precedence drift', 2, ORANGE)
para('The gateway accepts multiple candidate secrets (candidateSecrets) and only console.warns when a '
     'request matches a fallback instead of the registered secret. A stale/rotated secret keeps working '
     'silently; rotation is not actually enforced.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Make fallback-secret acceptance time-boxed and alarmed (metric + alert, not just a log line); '
          'fail closed once the rotation window ends.')

h('4.4 Session lifetime & cookie posture', 2, ORANGE)
para('Sessions last 12 hours (maxAge 60*60*12) with httpOnly, sameSite Lax, secure only when HTTPS is '
     'detected at runtime. No idle timeout, no rotation on privilege change.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Add a sliding idle timeout (30-60 min) with an absolute cap; force Secure in production regardless '
          'of runtime detection; rotate the session token on login and on role elevation.')

h('4.5 Dynamic SQL built by string concatenation', 2, ORANGE)
para('Several queries interpolate identifiers/clauses: SELECT COUNT(*) FROM "${t}", SELECT * FROM ${ds}, '
     'UPDATE users SET ${sets.join(", ")}, WHERE ${keyCol}=?. Values are parameterized, but table/column/'
     'dataset NAMES are concatenated. If any ${...} ever derives from user input, it is injectable.')
p = doc.add_paragraph(); p.add_run('Recommendation: ').bold = True
p.add_run('Whitelist every interpolated identifier against a fixed registry; never pass request-derived '
          'strings into an identifier position. Add a CI grep check for FROM ${ / SET ${.')

h('4.6 Randomness for receipts/scores', 2, YELLOW)
para('Math.random() is used for simulated receipts and the demo credit score (non-security). Session tokens '
     'correctly use crypto.randomUUID(). Recommendation: confirm all security-relevant randomness uses '
     'WebCrypto; add a lint fence so simulation helpers are never reused for secrets.')

h('4.7 Verbose operational logging', 2, YELLOW)
para('~25 console.* statements per app, some printing secret-precedence hints and client keys. Low risk but '
     'informative to an attacker with log access. Recommendation: route through a leveled logger; suppress '
     'debug/secret-hint logs in production.')

h('4.8 Satellite apps carry full host-gateway code (attack surface)', 2, YELLOW)
para('Because Feed/Nia ship a full 47-68 KB host gateway (finding 2.1), settlement/IPN endpoints may be '
     'reachable on satellites even though they should only act as tenants. Recommendation: ensure host-only '
     'routes are not mounted on satellites (or remove host code per 2.1); verify with a route inventory.')

# ---------------- 5. ROADMAP ----------------
h('5. Consolidation Roadmap (reduce complexity, resolve conflicts)', 1)
make_table(['Phase', 'Action', 'Resolves', 'Risk'],
    [['P1 (now)', 'Fix CORS allow-list; gate demo_otp behind DEMO_MODE; hash OTP at rest', '4.1, 4.2, 3.4', 'Low, additive'],
     ['P2', 'Extract shared pricing module + shared test vectors', '2.4, 2.7', 'Medium — parity tests'],
     ['P3', 'Unify payment gateway: one host (Equipment) + one client (satellites); delete host code from Feed/Nia', '2.1, 2.2, 2.5, 4.8', 'Medium/High — flag-gate'],
     ['P4', 'Canonicalize endpoints to /api/checkout/*; alias-redirect legacy; add missing Nia alias', '2.3', 'Low'],
     ['P5', 'Single resolveConfig() precedence contract; single appIdentity()', '2.6, 3.3', 'Low'],
     ['P6', 'One assertCheckoutAllowed() KYC guard; make cross-marketplace cash-only rule explicit', '3.1, 3.2', 'Low'],
     ['P7', 'Migrate privileged writes to runAsAdmin; standardize middleware dir', '2.9, 2.8', 'Low'],
     ['P8', 'Session hardening; SQL identifier whitelist; leveled logging', '4.4, 4.5, 4.7', 'Low']],
    widths=[0.8, 3.0, 1.3, 1.3])
para('')
p = doc.add_paragraph()
p.add_run('Guiding principle: ').bold = True
p.add_run('every phase is additive-first and flag-gated — the same zero-regression discipline already used '
          'for the offtake/checkout work. No phase removes a capability until its replacement is proven at '
          'parity in all three apps.')

# ---------------- 6. REFERENCE ----------------
h('6. Reference — App-identity dimensions (clarifies 3.3)', 1)
make_table(['Dimension', 'Where set', 'Meaning', 'Example values'],
    [['APP_TYPE', 'deploy env', 'Coarse app family for scope filtering', 'equipment, feed, input'],
     ['APP_MARKETPLACE', 'deploy env', 'Isolation floor for products a satellite may read', 'equipment, feed, mazao'],
     ['app_scope', 'product col', 'Which app family may surface this product', 'equipment, feed, both'],
     ['source_platform', 'product col', 'Which app originally listed the product', 'equipment, feed, mazao, merchant'],
     ['marketplace', 'product col', 'Storefront section / taxonomy tab', 'equipment, feeds, inputs']],
    widths=[1.3, 1.1, 2.4, 1.6])

doc.add_paragraph()
foot = doc.add_paragraph()
fr = foot.add_run('Generated by the Farmsky cross-app architecture & security review. Documentation only — '
                  'changes no runtime behaviour. Mirrored verbatim (SYSTEM_ARCHITECTURE_REVIEW.md) into all '
                  'three repositories so every team sees the same backlog.')
fr.italic = True; fr.font.size = Pt(9); fr.font.color.rgb = GREY

import sys
out = sys.argv[1] if len(sys.argv) > 1 else '/home/user/webapp/Farmsky_System_Review.docx'
doc.save(out)
print('saved', out)
